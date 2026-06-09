#!/usr/bin/env python3

from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_BASE_URL = "https://crm.controlitfactory.eu"
DEFAULT_WORKSPACE_ID = "63c3e463-fa79-4e0e-8f48-1363976393df"
DEFAULT_SOMIJA_PATH = "/Users/alexeykruminsh/Downloads/SOMIJA DCF CRM .xlsx"
DEFAULT_CRM_EXCEL_PATH = "/Users/alexeykruminsh/Downloads/CRM EXCEL .xlsx"
DEFAULT_BACKUP_DIR = ".codex/backups"
DEFAULT_SSH_HOST = "controlit-crm-vps"
DEFAULT_REMOTE_DIR = "/opt/controlit-crm"
DEFAULT_DB_SERVICE = "db"
DEFAULT_DB_USER = "controlit_user"
DEFAULT_DB_NAME = "default"
FIELD_WIDTH = 150

COUNTRY_OPTIONS = [
    ("Finland", "FINLAND", "blue"),
    ("Estonia", "ESTONIA", "sky"),
    ("Lithuania", "LITHUANIA", "green"),
    ("Latvia", "LATVIA", "turquoise"),
    ("United Arab Emirates", "UNITED_ARAB_EMIRATES", "purple"),
    ("Asia", "ASIA", "yellow"),
    ("Czechia", "CZECHIA", "orange"),
    ("Slovakia", "SLOVAKIA", "red"),
    ("Slovenia", "SLOVENIA", "pink"),
    ("Croatia", "CROATIA", "sky"),
    ("Romania", "ROMANIA", "green"),
    ("Hungary", "HUNGARY", "gray"),
    ("MENA region", "MENA", "purple"),
    ("Australia", "AUSTRALIA", "blue"),
    ("New Zealand", "NEW_ZEALAND", "green"),
]
COUNTRY_VALUE_ALIASES = {
    "uae": "UNITED_ARAB_EMIRATES",
    "u a e": "UNITED_ARAB_EMIRATES",
    "emirates": "UNITED_ARAB_EMIRATES",
    "czech republic": "CZECHIA",
    "mena": "MENA",
    "middle east and north africa": "MENA",
    "middle east north africa": "MENA",
    "united arab emirates uae": "UNITED_ARAB_EMIRATES",
    "uae united arab emirates": "UNITED_ARAB_EMIRATES",
    "new zealand": "NEW_ZEALAND",
    "nz": "NEW_ZEALAND",
}
PILOT_COMPANY_KEYS = {"yit", "sensor innovation", "dayone"}

EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.IGNORECASE)
COMPANY_SUFFIX_RE = re.compile(
    r"\b(oy|as|inc|ltd|llc|ab|group|gmbh|corp|corporation)\.?$",
    re.IGNORECASE,
)


COMPANY_FIELD_DEFINITIONS = [
    {
        "name": "companyCountry",
        "label": "Company country",
        "type": "SELECT",
        "icon": "IconMapPin",
        "description": "Country from CRM import source.",
        "options": [
            {"id": str(uuid.uuid4()), "label": label, "value": value, "color": color, "position": index}
            for index, (label, value, color) in enumerate(COUNTRY_OPTIONS)
        ],
    },
    {
        "name": "companyCategoryRaw",
        "label": "Company category raw",
        "type": "TEXT",
        "icon": "IconCategory",
        "description": "Original Category value from CRM import source.",
    },
    {
        "name": "projectTypes",
        "label": "Project types",
        "type": "TEXT",
        "icon": "IconBuilding",
        "description": "Original Project types value from CRM import source.",
    },
    {
        "name": "targetPersonRole",
        "label": "Target person/role",
        "type": "TEXT",
        "icon": "IconUserSearch",
        "description": "Original Target person/role to meet value from CRM import source.",
    },
]


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def collapse_spaces(value: str) -> str:
    return re.sub(r"\s+", " ", clean(value))


def normalize_company_key(value: str) -> str:
    normalized = collapse_spaces(value).casefold()
    normalized = re.sub(r"[’']", "", normalized)
    normalized = normalized.strip(" .,-")
    previous = None
    while previous != normalized:
        previous = normalized
        normalized = COMPANY_SUFFIX_RE.sub("", normalized).strip(" .,-")
    return normalized


def is_valid_company_name(value: str) -> bool:
    key = normalize_company_key(value)
    if key in {"", "-", "na", "n/a", "none", "null"}:
        return False
    if clean(value).startswith("'"):
        return False
    return True


def normalize_name_key(value: str) -> str:
    return collapse_spaces(value).casefold().strip(" .,-")


def normalize_email(value: str) -> str:
    match = EMAIL_RE.search(clean(value))
    return match.group(0).lower() if match else ""


def normalize_phone(value: str) -> str:
    phone = clean(value)
    if phone.startswith("'"):
        phone = phone[1:]
    return collapse_spaces(phone)


def normalize_linkedin(value: str) -> tuple[str, str]:
    link = clean(value)
    if not link:
        return "", ""
    lower = link.lower()
    if lower.startswith("http://") or lower.startswith("https://"):
        return link, link
    if lower.startswith("linkedin.com/") or lower.startswith("www.linkedin.com/"):
        normalized = f"https://{link}"
        return normalized, normalized
    return "", link


def normalize_website_for_storage(value: str) -> str:
    website = clean(value)
    if not website:
        return ""
    website = website.split()[0].strip(" ,;")
    if website.lower().startswith(("http://", "https://")):
        normalized = website
    else:
        normalized = f"https://{website}"
    return normalized if is_valid_http_url(normalized) else ""


def is_valid_http_url(value: str) -> bool:
    parsed = urllib.parse.urlparse(clean(value))
    if parsed.scheme.lower() not in {"http", "https"}:
        return False
    host = parsed.netloc.strip().lower()
    if not host:
        return False
    if any(character in host for character in "'\"<>[]{}|\\^"):
        return False
    if host in {"-", "na", "n/a", "none", "null"}:
        return False
    return "." in host or host == "localhost"


def normalize_domain_key(value: str) -> str:
    website = clean(value)
    if not website:
        return ""
    if "://" not in website:
        website_for_parse = f"https://{website}"
    else:
        website_for_parse = website
    parsed = urllib.parse.urlparse(website_for_parse)
    host = (parsed.netloc or parsed.path).lower()
    host = host.split("/")[0].strip()
    if host.startswith("www."):
        host = host[4:]
    return host


def link_label(value: str) -> str:
    domain = normalize_domain_key(value)
    return domain or clean(value)


def normalize_country_key(country: str) -> str:
    normalized = clean(country).casefold()
    normalized = normalized.replace("&", " and ")
    normalized = re.sub(r"[/(),.;:_-]+", " ", normalized)
    normalized = re.sub(r"\s+", " ", normalized)
    return normalized.strip()


def country_value(country: str) -> str | None:
    normalized = normalize_country_key(country)
    values_by_key = {
        **{normalize_country_key(label): value for label, value, _color in COUNTRY_OPTIONS},
        **{normalize_country_key(value): value for _label, value, _color in COUNTRY_OPTIONS},
        **COUNTRY_VALUE_ALIASES,
    }
    return values_by_key.get(normalized)


def map_company_type(category: str) -> str | None:
    text = clean(category).casefold()
    if not text:
        return None
    if "distributor" in text or "distribution" in text:
        return "DISTIBUTORS"
    if "architect" in text or "architecture" in text or "design studio" in text:
        return "ARCHITECTURAL_BUREAU"
    if "construction" in text or "contractor" in text or "builder" in text:
        return "BUILDERS"
    if "developer" in text or "development" in text or "property" in text or "real estate" in text:
        return "DEVELOPER"
    return None


def first_non_empty(*values: str) -> str:
    for value in values:
        cleaned = clean(value)
        if cleaned:
            return cleaned
    return ""


def blocknote_body(markdown: str) -> str:
    lines = [line for line in markdown.splitlines() if line.strip()]
    blocks = []
    for index, line in enumerate(lines or [""]):
        blocks.append(
            {
                "id": f"import-{index}-{hashlib.sha1(line.encode('utf-8')).hexdigest()[:8]}",
                "type": "paragraph",
                "props": {
                    "textColor": "default",
                    "backgroundColor": "default",
                    "textAlignment": "left",
                },
                "content": [{"type": "text", "text": line, "styles": {}}],
                "children": [],
            }
        )
    return json.dumps(blocks, ensure_ascii=False)


def import_marker(kind: str, key: str, body: str) -> str:
    digest = hashlib.sha1(f"{kind}|{key}|{body}".encode("utf-8")).hexdigest()[:16]
    return f"CRM_IMPORT:{kind}:{digest}"


def parse_source_records(somija_path: str, crm_excel_path: str) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    records.extend(parse_somija_records(somija_path))
    records.extend(parse_crm_excel_records(crm_excel_path))
    return records


def parse_somija_records(path: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    header_row = None
    headers: list[str] = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = [clean(value) for value in row]
        if "Company" in values and "Email" in values:
            header_row = row_number
            headers = values
            break
    if header_row is None:
        raise RuntimeError(f"Could not find header row in {path}")

    records = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        values = {
            headers[index]: clean(row[index]) if index < len(row) else ""
            for index in range(len(headers))
            if headers[index]
        }
        if not any(values.values()):
            continue
        linkedin_url, linkedin_note = normalize_linkedin(values.get("LinkedIn", ""))
        first_name = clean(values.get("First name", ""))
        last_name = clean(values.get("Last name", ""))
        records.append(
            {
                "source_file": Path(path).name,
                "source_sheet": sheet.title,
                "source_row": row_number,
                "country": "Finland",
                "company": clean(values.get("Company", "")),
                "company_category": "",
                "company_website": "",
                "company_city": "",
                "project_types": "",
                "target_role_to_meet": "",
                "company_notes": "",
                "contact_full_name": collapse_spaces(f"{first_name} {last_name}"),
                "contact_first_name": first_name,
                "contact_last_name": last_name,
                "contact_position": clean(values.get("Role", "")),
                "contact_email": clean(values.get("Email", "")),
                "contact_phone": normalize_phone(values.get("Phone", "")),
                "contact_linkedin": linkedin_url,
                "contact_note": linkedin_note,
            }
        )
    return records


def parse_crm_excel_records(path: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    records = []
    for sheet in workbook.worksheets:
        rows = [[clean(value) for value in row] for row in sheet.iter_rows(values_only=True)]
        country = sheet.title.replace("CONTACTS", "").strip().title()
        index = 0
        while index < len(rows):
            first_cell = rows[index][0] if rows[index] else ""
            if first_cell.startswith("Company #") or first_cell == "Company Name":
                company = {
                    "source_file": Path(path).name,
                    "source_sheet": sheet.title,
                    "country": country,
                    "company": "",
                    "company_category": "",
                    "company_website": "",
                    "company_city": "",
                    "project_types": "",
                    "target_role_to_meet": "",
                    "company_notes": "",
                }
                if first_cell.startswith("Company #"):
                    index += 1

                while index < len(rows):
                    cell_a, cell_b, cell_c, cell_d = (rows[index] + ["", "", "", ""])[:4]
                    if cell_a.startswith("Company #") and company["company"]:
                        index -= 1
                        break
                    if cell_a == "Contacts":
                        break
                    for key, value in ((cell_a, cell_b), (cell_c, cell_d)):
                        normalized_key = key.strip().lower()
                        if normalized_key == "company name":
                            company["company"] = value
                        elif normalized_key == "category":
                            company["company_category"] = value
                        elif normalized_key == "website":
                            company["company_website"] = value
                        elif normalized_key == "city":
                            company["company_city"] = value
                        elif normalized_key.startswith("project types"):
                            company["project_types"] = value
                        elif normalized_key.startswith("target person"):
                            company["target_role_to_meet"] = value
                        elif normalized_key == "notes":
                            company["company_notes"] = value
                    index += 1

                while index < len(rows) and rows[index][0] != "Name":
                    if rows[index][0].startswith("Company #"):
                        index -= 1
                        break
                    index += 1

                if index < len(rows) and rows[index][0] == "Name":
                    index += 1
                    while index < len(rows):
                        cell_a, cell_b, cell_c, cell_d = (rows[index] + ["", "", "", ""])[:4]
                        if cell_a.startswith("Company #") or cell_a == "Company Name":
                            index -= 1
                            break
                        if any([cell_a, cell_b, cell_c, cell_d]):
                            first_name, last_name = split_full_name(cell_a)
                            records.append(
                                {
                                    **company,
                                    "source_row": index + 1,
                                    "contact_full_name": cell_a,
                                    "contact_first_name": first_name,
                                    "contact_last_name": last_name,
                                    "contact_position": cell_b,
                                    "contact_email": cell_c,
                                    "contact_phone": normalize_phone(cell_d),
                                    "contact_linkedin": "",
                                    "contact_note": "",
                                }
                            )
                        index += 1
            index += 1
    return records


def split_full_name(full_name: str) -> tuple[str, str]:
    parts = collapse_spaces(full_name).split()
    if not parts:
        return "", ""
    return parts[0], " ".join(parts[1:])


def merge_source_companies(records: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    companies: dict[str, dict[str, Any]] = {}
    domain_to_company_key: dict[str, str] = {}
    for record in records:
        company_name = clean(record.get("company", ""))
        if not is_valid_company_name(company_name):
            continue
        company_key = normalize_company_key(company_name)
        domain_key = normalize_domain_key(normalize_website_for_storage(record.get("company_website", "")))
        merge_key = domain_to_company_key.get(domain_key) if domain_key else None
        if not merge_key:
            merge_key = company_key
        company = companies.setdefault(
            merge_key,
            {
                "key": merge_key,
                "name": company_name,
                "country": "",
                "category": "",
                "website": "",
                "city": "",
                "project_types": "",
                "target_person_role": "",
                "company_notes": [],
                "contacts": [],
                "source_rows": [],
            },
        )
        if domain_key:
            domain_to_company_key.setdefault(domain_key, merge_key)
        company["country"] = first_non_empty(company["country"], record.get("country", ""))
        company["category"] = first_non_empty(company["category"], record.get("company_category", ""))
        company["website"] = first_non_empty(company["website"], record.get("company_website", ""))
        company["city"] = first_non_empty(company["city"], record.get("company_city", ""))
        company["project_types"] = first_non_empty(company["project_types"], record.get("project_types", ""))
        company["target_person_role"] = first_non_empty(company["target_person_role"], record.get("target_role_to_meet", ""))
        note = clean(record.get("company_notes", ""))
        if note and note not in company["company_notes"]:
            company["company_notes"].append(note)
        company["contacts"].append(record)
        company["source_rows"].append(f"{record.get('source_file')} / {record.get('source_sheet')} row {record.get('source_row')}")
    return companies


def existing_company_indexes(companies: list[dict[str, Any]]) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    by_domain = {}
    by_name = {}
    for company in companies:
        domain_key = normalize_domain_key(link_primary_url(company, "domainName"))
        name_key = normalize_company_key(company.get("name", ""))
        if domain_key:
            by_domain.setdefault(domain_key, company)
        if name_key:
            by_name.setdefault(name_key, company)
    return by_domain, by_name


def existing_person_indexes(people: list[dict[str, Any]]) -> tuple[dict[str, dict[str, Any]], dict[tuple[str, str], dict[str, Any]]]:
    by_email = {}
    by_company_name = {}
    for person in people:
        email = normalize_email(email_primary(person))
        if email:
            by_email.setdefault(email, person)
        first_name, last_name = person_name_parts(person)
        full_name = collapse_spaces(f"{first_name} {last_name}")
        name_key = normalize_name_key(full_name)
        company_id = clean(person.get("companyId", ""))
        if company_id and name_key:
            by_company_name.setdefault((company_id, name_key), person)
    return by_email, by_company_name


COMPOSITE_LEGACY_KEYS = {
    "domainName": {
        "primaryLinkUrl": "domainNamePrimaryLinkUrl",
        "primaryLinkLabel": "domainNamePrimaryLinkLabel",
    },
    "address": {
        "addressCity": "addressAddressCity",
        "addressCountry": "addressAddressCountry",
    },
    "name": {
        "firstName": "nameFirstName",
        "lastName": "nameLastName",
    },
    "emails": {
        "primaryEmail": "emailsPrimaryEmail",
    },
    "phones": {
        "primaryPhoneNumber": "phonesPrimaryPhoneNumber",
    },
    "linkedinLink": {
        "primaryLinkUrl": "linkedinLinkPrimaryLinkUrl",
        "primaryLinkLabel": "linkedinLinkPrimaryLinkLabel",
    },
}


def is_empty_value(value: Any) -> bool:
    if value is None or value == "":
        return True
    if isinstance(value, list):
        return len(value) == 0
    if isinstance(value, dict):
        return all(is_empty_value(nested_value) for nested_value in value.values())
    return False


def compact_payload(payload: dict[str, Any]) -> dict[str, Any]:
    compacted: dict[str, Any] = {}
    for key, value in payload.items():
        if isinstance(value, dict):
            nested = {nested_key: nested_value for nested_key, nested_value in value.items() if not is_empty_value(nested_value)}
            if nested:
                compacted[key] = nested
        elif not is_empty_value(value):
            compacted[key] = value
    return compacted


def existing_field_value(existing: dict[str, Any], key: str) -> Any:
    value = existing.get(key)
    if not is_empty_value(value) or key in existing:
        return value
    if key not in COMPOSITE_LEGACY_KEYS:
        return value
    return {
        nested_key: existing.get(legacy_key)
        for nested_key, legacy_key in COMPOSITE_LEGACY_KEYS[key].items()
    }


def existing_nested_value(existing: dict[str, Any], key: str, nested_key: str) -> Any:
    value = existing.get(key)
    if isinstance(value, dict):
        return value.get(nested_key)
    legacy_key = COMPOSITE_LEGACY_KEYS.get(key, {}).get(nested_key)
    if legacy_key:
        return existing.get(legacy_key)
    return None


def link_primary_url(record: dict[str, Any], key: str) -> str:
    return clean(existing_nested_value(record, key, "primaryLinkUrl"))


def email_primary(person: dict[str, Any]) -> str:
    return clean(existing_nested_value(person, "emails", "primaryEmail"))


def person_name_parts(person: dict[str, Any]) -> tuple[str, str]:
    return (
        clean(existing_nested_value(person, "name", "firstName")),
        clean(existing_nested_value(person, "name", "lastName")),
    )


def source_company_payload(company: dict[str, Any]) -> dict[str, Any]:
    website = normalize_website_for_storage(company.get("website", ""))
    payload = {
        "name": clean(company.get("name", "")),
        "domainName": {
            "primaryLinkUrl": website,
            "primaryLinkLabel": link_label(website),
        },
        "address": {
            "addressCity": clean(company.get("city", "")),
            "addressCountry": clean(company.get("country", "")),
        },
        "companyCountry": country_value(company.get("country", "")),
        "companyCategoryRaw": clean(company.get("category", "")),
        "projectTypes": clean(company.get("project_types", "")),
        "targetPersonRole": clean(company.get("target_person_role", "")),
        "companyType": map_company_type(company.get("category", "")),
    }
    return compact_payload(payload)


def fill_only_update(existing: dict[str, Any], desired: dict[str, Any]) -> dict[str, Any]:
    update = {}
    for key, value in desired.items():
        if is_empty_value(value):
            continue
        if isinstance(value, dict):
            nested_update = {
                nested_key: nested_value
                for nested_key, nested_value in value.items()
                if not is_empty_value(nested_value)
                and is_empty_value(existing_nested_value(existing, key, nested_key))
            }
            if nested_update:
                update[key] = nested_update
        elif is_empty_value(existing_field_value(existing, key)):
            update[key] = value
    return update


def merge_payload_missing(target: dict[str, Any], source: dict[str, Any]) -> None:
    for key, value in source.items():
        if is_empty_value(value):
            continue
        if isinstance(value, dict):
            target_value = target.setdefault(key, {})
            if not isinstance(target_value, dict):
                continue
            for nested_key, nested_value in value.items():
                if not is_empty_value(nested_value) and is_empty_value(target_value.get(nested_key)):
                    target_value[nested_key] = nested_value
        elif is_empty_value(target.get(key)):
            target[key] = value


def contact_payload(contact: dict[str, Any], company_id: str | None) -> dict[str, Any]:
    first_name = clean(contact.get("contact_first_name", ""))
    last_name = clean(contact.get("contact_last_name", ""))
    if not first_name and not last_name:
        first_name, last_name = split_full_name(contact.get("contact_full_name", ""))
    phone = normalize_phone(contact.get("contact_phone", ""))
    phone_payload = {"primaryPhoneNumber": phone}
    if phone and not phone.startswith("+"):
        phone_payload.update(phone_metadata_for_country(contact.get("country", "")))
    linkedin = clean(contact.get("contact_linkedin", ""))
    payload = {
        "name": {
            "firstName": first_name,
            "lastName": last_name,
        },
        "emails": {
            "primaryEmail": normalize_email(contact.get("contact_email", "")),
        },
        "phones": phone_payload,
        "jobTitle": clean(contact.get("contact_position", "")),
        "linkedinLink": {
            "primaryLinkUrl": linkedin,
            "primaryLinkLabel": linkedin,
        },
        "companyId": company_id,
        "personTerritory": country_value(contact.get("country", "")),
    }
    return compact_payload(payload)


def phone_metadata_for_country(country: str) -> dict[str, str]:
    normalized = clean(country).casefold()
    if normalized == "finland":
        return {"primaryPhoneCountryCode": "FI", "primaryPhoneCallingCode": "358"}
    if normalized == "estonia":
        return {"primaryPhoneCountryCode": "EE", "primaryPhoneCallingCode": "372"}
    if normalized == "lithuania":
        return {"primaryPhoneCountryCode": "LT", "primaryPhoneCallingCode": "370"}
    return {}


def person_aliases(company_key: str, full_name: str, email: str) -> list[str]:
    aliases = []
    if email:
        aliases.append(f"email:{email}")
    name_key = normalize_name_key(full_name)
    if company_key and name_key:
        aliases.append(f"name:{company_key}|{name_key}")
    return aliases


def first_planned_person_entry(
    planned_person_aliases: dict[str, dict[str, Any]],
    aliases: list[str],
) -> dict[str, Any] | None:
    for alias in aliases:
        if alias in planned_person_aliases:
            return planned_person_aliases[alias]
    return None


def register_person_aliases(
    planned_person_aliases: dict[str, dict[str, Any]],
    aliases: list[str],
    entry: dict[str, Any],
) -> None:
    for alias in aliases:
        planned_person_aliases[alias] = entry


def build_import_plan(source_records: list[dict[str, Any]], existing: dict[str, Any]) -> dict[str, Any]:
    source_companies = merge_source_companies(source_records)
    existing_companies = existing.get("companies", [])
    existing_people = existing.get("people", [])
    existing_note_markers = existing.get("note_import_keys", set())
    by_domain, by_name = existing_company_indexes(existing_companies)
    people_by_email, people_by_company_name = existing_person_indexes(existing_people)

    plan: dict[str, Any] = {
        "company_creates": [],
        "company_updates": [],
        "person_creates": [],
        "person_updates": [],
        "company_notes": [],
        "skipped_person_rows": [],
        "skipped_company_rows": [],
        "source_duplicate_emails": {},
        "source_company_count": len(source_companies),
        "company_ids_by_key": {},
    }
    planned_company_ids: dict[str, str | None] = {}
    planned_person_aliases: dict[str, dict[str, Any]] = {}
    source_email_counts: dict[str, int] = {}

    for contact in source_records:
        if not is_valid_company_name(contact.get("company", "")):
            plan["skipped_company_rows"].append(source_line(contact))
            continue
        email = normalize_email(contact.get("contact_email", ""))
        if email:
            source_email_counts[email] = source_email_counts.get(email, 0) + 1
    plan["source_duplicate_emails"] = {
        email: count for email, count in sorted(source_email_counts.items()) if count > 1
    }

    for company_key, company in source_companies.items():
        desired = source_company_payload(company)
        domain_key = normalize_domain_key(link_primary_url(desired, "domainName"))
        existing_company = (by_domain.get(domain_key) if domain_key else None) or by_name.get(company_key)
        if existing_company:
            planned_company_ids[company_key] = existing_company["id"]
            plan["company_ids_by_key"][company_key] = existing_company["id"]
            update = fill_only_update(existing_company, desired)
            if update:
                plan["company_updates"].append({"id": existing_company["id"], "company_key": company_key, "data": update})
        else:
            planned_company_ids[company_key] = None
            plan["company_creates"].append({"company_key": company_key, "data": desired})

        for note_body in company.get("company_notes", []):
            add_company_note(
                plan,
                existing_note_markers,
                company_key,
                "Company source notes",
                note_body,
                country_value(company.get("country", "")),
            )

        for contact in company["contacts"]:
            full_name = clean(contact.get("contact_full_name", ""))
            email = normalize_email(contact.get("contact_email", ""))
            phone = normalize_phone(contact.get("contact_phone", ""))
            position = clean(contact.get("contact_position", ""))
            contact_note = clean(contact.get("contact_note", ""))

            if not full_name:
                if email or phone or position:
                    generic_lines = ["Generic contact from source:"]
                    if position:
                        generic_lines.append(f"Role: {position}")
                    if email:
                        generic_lines.append(f"Email: {email}")
                    if phone:
                        generic_lines.append(f"Phone: {phone}")
                    generic_lines.append(source_line(contact))
                    add_company_note(
                        plan,
                        existing_note_markers,
                        company_key,
                        "Generic contact",
                        "\n".join(generic_lines),
                        country_value(contact.get("country", "")),
                    )
                plan["skipped_person_rows"].append(source_line(contact))
                continue

            if contact_note:
                add_company_note(
                    plan,
                    existing_note_markers,
                    company_key,
                    "Contact note",
                    f"{full_name}: {contact_note}\n{source_line(contact)}",
                    country_value(contact.get("country", "")),
                )

            company_id = planned_company_ids.get(company_key)
            aliases = person_aliases(company_key, full_name, email)
            payload = contact_payload(contact, company_id)
            planned_entry = first_planned_person_entry(planned_person_aliases, aliases)
            if planned_entry:
                if planned_entry["kind"] == "create":
                    merge_payload_missing(planned_entry["action"]["data"], payload)
                else:
                    delta = fill_only_update(planned_entry["existing"], payload)
                    if delta:
                        if planned_entry["kind"] == "existing":
                            action = {
                                "id": planned_entry["existing"]["id"],
                                "company_key": company_key,
                                "data": delta,
                            }
                            plan["person_updates"].append(action)
                            planned_entry["kind"] = "update"
                            planned_entry["action"] = action
                        else:
                            merge_payload_missing(planned_entry["action"]["data"], delta)
                register_person_aliases(planned_person_aliases, aliases, planned_entry)
                continue

            existing_person = people_by_email.get(email) if email else None
            if not existing_person and company_id:
                existing_person = people_by_company_name.get((company_id, normalize_name_key(full_name)))

            if existing_person:
                payload = contact_payload(contact, company_id or existing_person.get("companyId"))
                update = fill_only_update(existing_person, payload)
                entry = {"kind": "existing", "existing": existing_person}
                if update:
                    action = {"id": existing_person["id"], "company_key": company_key, "data": update}
                    plan["person_updates"].append(action)
                    entry = {"kind": "update", "existing": existing_person, "action": action}
                register_person_aliases(planned_person_aliases, aliases, entry)
            else:
                action = {"company_key": company_key, "person_key": aliases[0], "data": payload}
                plan["person_creates"].append(action)
                register_person_aliases(
                    planned_person_aliases,
                    aliases,
                    {"kind": "create", "action": action},
                )

    return plan


def add_company_note(
    plan: dict[str, Any],
    existing_markers: set[str],
    company_key: str,
    title: str,
    body: str,
    territory: str | None,
) -> None:
    marker = import_marker("company-note", company_key, body)
    if marker in existing_markers:
        return
    plan["company_notes"].append(
        {
            "company_key": company_key,
            "title": f"CRM import - {title}",
            "body": f"{body}\n\n[{marker}]",
            "marker": marker,
            "territory": territory,
        }
    )


def source_line(contact: dict[str, Any]) -> str:
    return f"Source: {contact.get('source_file')} / {contact.get('source_sheet')} row {contact.get('source_row')}"


@dataclass
class CrmConfig:
    base_url: str
    admin_email: str
    admin_password: str
    workspace_id: str
    backup_dir: Path
    ssh_host: str
    remote_dir: str
    db_service: str
    db_user: str
    db_name: str


class CrmClient:
    def __init__(self, config: CrmConfig):
        self.config = config
        self.workspace: dict[str, Any] | None = None
        self.token = ""

    def authenticate(self) -> None:
        sign_in = self.metadata_request(
            """
            mutation SignIn($email: String!, $password: String!) {
              signIn(email: $email, password: $password) {
                availableWorkspaces {
                  availableWorkspacesForSignIn {
                    id
                    displayName
                    loginToken
                    workspaceUrls {
                      customUrl
                      subdomainUrl
                    }
                  }
                }
              }
            }
            """,
            {"email": self.config.admin_email, "password": self.config.admin_password},
            token=None,
        )
        workspaces = sign_in["signIn"]["availableWorkspaces"]["availableWorkspacesForSignIn"]
        self.workspace = next((workspace for workspace in workspaces if workspace["id"] == self.config.workspace_id), None)
        if not self.workspace:
            raise RuntimeError(f"Workspace not found: {self.config.workspace_id}")
        origin = (
            self.workspace["workspaceUrls"].get("customUrl")
            or self.workspace["workspaceUrls"].get("subdomainUrl")
            or self.config.base_url
        )
        auth = self.metadata_request(
            """
            mutation GetAuthTokensFromLoginToken($loginToken: String!, $origin: String!) {
              getAuthTokensFromLoginToken(loginToken: $loginToken, origin: $origin) {
                tokens {
                  accessOrWorkspaceAgnosticToken {
                    token
                  }
                }
              }
            }
            """,
            {"loginToken": self.workspace["loginToken"], "origin": origin},
            token=None,
        )
        self.token = auth["getAuthTokensFromLoginToken"]["tokens"]["accessOrWorkspaceAgnosticToken"]["token"]

    def metadata_request(self, query: str, variables: dict[str, Any] | None = None, token: str | None = "USE_CURRENT") -> Any:
        headers = {"Content-Type": "application/json"}
        active_token = self.token if token == "USE_CURRENT" else token
        if active_token:
            headers["Authorization"] = f"Bearer {active_token}"
        return self.request_json(
            "POST",
            f"{self.config.base_url}/metadata",
            {"query": query, "variables": variables or {}},
            headers,
        )["data"]

    def rest(self, method: str, path: str, body: dict[str, Any] | None = None) -> Any:
        return self.request_json(
            method,
            f"{self.config.base_url}{path}",
            body,
            {"Content-Type": "application/json", "Authorization": f"Bearer {self.token}"},
        )

    @staticmethod
    def request_json(method: str, url: str, body: dict[str, Any] | None, headers: dict[str, str]) -> Any:
        data = json.dumps(body).encode("utf-8") if body is not None else None
        request = urllib.request.Request(url, data=data, headers=headers, method=method)
        try:
            with urllib.request.urlopen(request, timeout=60) as response:
                payload = response.read().decode("utf-8")
                parsed = json.loads(payload) if payload else None
        except urllib.error.HTTPError as error:
            payload = error.read().decode("utf-8", errors="replace")
            raise RuntimeError(f"{method} {url} failed: {error.code} {payload}") from error
        if isinstance(parsed, dict) and parsed.get("errors"):
            raise RuntimeError(f"{method} {url} failed: {json.dumps(parsed['errors'], ensure_ascii=False)}")
        return parsed


def create_database_backup(config: CrmConfig, prefix: str) -> Path:
    config.backup_dir.mkdir(parents=True, exist_ok=True)
    timestamp = time.strftime("%Y%m%d-%H%M%S")
    backup_path = config.backup_dir / f"{prefix}-{timestamp}.sql.gz"
    remote_command = (
        f"cd {shell_quote(config.remote_dir)} && "
        f"docker compose exec -T {shell_quote(config.db_service)} "
        f"pg_dump -U {shell_quote(config.db_user)} {shell_quote(config.db_name)}"
    )
    with subprocess.Popen(
        ["ssh", config.ssh_host, remote_command],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    ) as process:
        with gzip.open(backup_path, "wb") as output:
            shutil.copyfileobj(process.stdout, output)
        _, stderr = process.communicate()
        if process.returncode != 0:
            backup_path.unlink(missing_ok=True)
            raise RuntimeError(f"Database backup failed: {stderr.decode('utf-8', errors='replace')}")
    return backup_path


def shell_quote(value: str) -> str:
    return "'" + value.replace("'", "'\\''") + "'"


def get_company_metadata(client: CrmClient) -> dict[str, Any]:
    data = client.metadata_request(
        """
        query ObjectMetadataItems {
          objects(paging: { first: 1000 }) {
            edges {
              node {
                id
                nameSingular
                namePlural
                labelPlural
                fieldsList {
                  id
                  name
                  label
                  type
                  options
                }
              }
            }
          }
        }
        """
    )
    objects = [edge["node"] for edge in data["objects"]["edges"]]
    company = next((obj for obj in objects if obj["nameSingular"] == "company"), None)
    if not company:
        raise RuntimeError("Company object metadata not found")
    return company


def ensure_company_metadata(client: CrmClient, apply: bool) -> list[str]:
    company = get_company_metadata(client)
    fields_by_name = {field["name"]: field for field in company["fieldsList"]}
    actions: list[str] = []
    for definition in COMPANY_FIELD_DEFINITIONS:
        existing = fields_by_name.get(definition["name"])
        if existing:
            if existing["type"] != definition["type"]:
                raise RuntimeError(
                    f"Field {definition['name']} exists as {existing['type']}, expected {definition['type']}"
                )
            if definition.get("options") and normalize_options(existing.get("options") or []) != normalize_options(definition["options"]):
                raise RuntimeError(f"Field {definition['name']} exists with different options")
            continue
        actions.append(f"create field {definition['name']}")
        if apply:
            created = client.metadata_request(
                """
                mutation CreateOneFieldMetadataItem($input: CreateOneFieldMetadataInput!) {
                  createOneField(input: $input) {
                    id
                    name
                    label
                    type
                    options
                  }
                }
                """,
                {
                    "input": {
                        "field": {
                            "objectMetadataId": company["id"],
                            "name": definition["name"],
                            "label": definition["label"],
                            "type": definition["type"],
                            "icon": definition["icon"],
                            "description": definition["description"],
                            "isCustom": True,
                            "isActive": True,
                            "isNullable": True,
                            "isLabelSyncedWithName": False,
                            **({"options": definition["options"]} if definition.get("options") else {}),
                        }
                    }
                },
            )["createOneField"]
            fields_by_name[created["name"]] = created

    if not apply and any(action.startswith("create field ") for action in actions):
        return actions

    refreshed_company = get_company_metadata(client) if apply and actions else company
    fields_by_name = {field["name"]: field for field in refreshed_company["fieldsList"]}
    views = client.rest("GET", f"/rest/metadata/views?objectMetadataId={urllib.parse.quote(refreshed_company['id'])}")
    all_companies_view = next((view for view in views if view["type"] == "TABLE" and view["position"] == 0), None)
    if not all_companies_view:
        raise RuntimeError("All Companies table view not found")
    view_fields = client.rest("GET", f"/rest/metadata/viewFields?viewId={urllib.parse.quote(all_companies_view['id'])}")
    view_fields_by_field_id = {view_field["fieldMetadataId"]: view_field for view_field in view_fields}
    company_type = fields_by_name.get("companyType")
    base_position = max(float(view_field["position"]) for view_field in view_fields) if view_fields else 0
    import_view_positions = [
        float(view_fields_by_field_id[fields_by_name[definition["name"]]["id"]]["position"])
        for definition in COMPANY_FIELD_DEFINITIONS
        if fields_by_name[definition["name"]]["id"] in view_fields_by_field_id
    ]
    if import_view_positions:
        base_position = min(import_view_positions) - 1
    elif company_type and company_type["id"] in view_fields_by_field_id:
        base_position = float(view_fields_by_field_id[company_type["id"]]["position"])
    for offset, field_name in enumerate([definition["name"] for definition in COMPANY_FIELD_DEFINITIONS], start=1):
        field = fields_by_name[field_name]
        existing_view_field = view_fields_by_field_id.get(field["id"])
        position = base_position + offset
        if not existing_view_field:
            actions.append(f"add {field_name} to All Companies view")
            if apply:
                client.rest(
                    "POST",
                    "/rest/metadata/viewFields",
                    {
                        "fieldMetadataId": field["id"],
                        "viewId": all_companies_view["id"],
                        "isVisible": True,
                        "size": FIELD_WIDTH,
                        "position": position,
                    },
                )
        elif (
            existing_view_field.get("isVisible") is not True
            or float(existing_view_field.get("position", 0)) != position
            or int(existing_view_field.get("size", 0)) != FIELD_WIDTH
        ):
            actions.append(f"update {field_name} in All Companies view")
            if apply:
                client.rest(
                    "PATCH",
                    f"/rest/metadata/viewFields/{urllib.parse.quote(existing_view_field['id'])}",
                    {"isVisible": True, "size": FIELD_WIDTH, "position": position},
                )
    return actions


def normalize_options(options: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(
        [
            {
                "label": option.get("label"),
                "value": option.get("value"),
                "color": option.get("color"),
                "position": int(option.get("position", 0)),
            }
            for option in options
        ],
        key=lambda option: option["position"],
    )


def fetch_all(client: CrmClient, object_name_plural: str) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    after = ""
    while True:
        query = {"limit": "1000"}
        if after:
            query["starting_after"] = after
        path = f"/rest/{object_name_plural}?{urllib.parse.urlencode(query)}"
        response = client.rest("GET", path)
        batch = response["data"].get(object_name_plural, [])
        records.extend(batch)
        page_info = response.get("pageInfo") or {}
        if not page_info.get("hasNextPage"):
            break
        after = page_info.get("endCursor")
        if not after:
            break
    return records


def fetch_existing_state(client: CrmClient) -> dict[str, Any]:
    companies = fetch_all(client, "companies")
    people = fetch_all(client, "people")
    notes = fetch_all(client, "notes")
    markers = set()
    for note in notes:
        body_v2 = note.get("bodyV2") if isinstance(note.get("bodyV2"), dict) else {}
        for text in (
            note.get("bodyV2Markdown", ""),
            body_v2.get("markdown", ""),
            note.get("title", ""),
        ):
            markers.update(re.findall(r"CRM_IMPORT:[^\]\s]+", clean(text)))
    return {"companies": companies, "people": people, "note_import_keys": markers}


def filter_records(records: list[dict[str, Any]], pilot: bool, limit_companies: int | None) -> list[dict[str, Any]]:
    if pilot:
        return [record for record in records if normalize_company_key(record.get("company", "")) in PILOT_COMPANY_KEYS]
    if limit_companies:
        selected_keys = []
        seen = set()
        for record in records:
            key = normalize_company_key(record.get("company", ""))
            if key and key not in seen:
                seen.add(key)
                selected_keys.append(key)
            if len(selected_keys) >= limit_companies:
                break
        selected = set(selected_keys)
        return [record for record in records if normalize_company_key(record.get("company", "")) in selected]
    return records


def apply_import_plan(client: CrmClient, plan: dict[str, Any], existing_state: dict[str, Any]) -> dict[str, int]:
    counts = {
        "companies_created": 0,
        "companies_updated": 0,
        "people_created": 0,
        "people_updated": 0,
        "notes_created": 0,
    }
    company_id_by_key = {**plan.get("company_ids_by_key", {})}
    for company in existing_state["companies"]:
        key = normalize_company_key(company.get("name", ""))
        if key:
            company_id_by_key.setdefault(key, company["id"])

    for action in plan["company_creates"]:
        created = client.rest("POST", "/rest/companies", action["data"])["data"]["createCompany"]
        company_id_by_key[action["company_key"]] = created["id"]
        counts["companies_created"] += 1

    for action in plan["company_updates"]:
        client.rest("PATCH", f"/rest/companies/{urllib.parse.quote(action['id'])}", action["data"])
        company_id_by_key[action["company_key"]] = action["id"]
        counts["companies_updated"] += 1

    for action in plan["person_creates"]:
        company_id = company_id_by_key.get(action["company_key"])
        data = {**action["data"]}
        if company_id and not data.get("companyId"):
            data["companyId"] = company_id
        client.rest("POST", "/rest/people", data)
        counts["people_created"] += 1

    for action in plan["person_updates"]:
        data = {**action["data"]}
        company_id = company_id_by_key.get(action["company_key"])
        if company_id and not data.get("companyId"):
            data["companyId"] = company_id
        client.rest("PATCH", f"/rest/people/{urllib.parse.quote(action['id'])}", data)
        counts["people_updated"] += 1

    for note in plan["company_notes"]:
        company_id = company_id_by_key.get(note["company_key"])
        if not company_id:
            raise RuntimeError(f"Cannot create note; company id missing for {note['company_key']}")
        created_note = client.rest(
            "POST",
            "/rest/notes",
            note_payload(note),
        )["data"]["createNote"]
        client.rest("POST", "/rest/noteTargets", {"noteId": created_note["id"], "companyId": company_id})
        counts["notes_created"] += 1

    return counts


def note_payload(note: dict[str, Any]) -> dict[str, Any]:
    payload = {
        "title": note["title"],
        "bodyV2": {
            "markdown": note["body"],
        },
    }
    territory = country_value(note.get("territory", ""))
    if territory:
        payload["noteTerritory"] = territory
    return payload


def summarize_plan(plan: dict[str, Any]) -> dict[str, Any]:
    return {
        "source_companies": plan["source_company_count"],
        "companies_to_create": len(plan["company_creates"]),
        "companies_to_update": len(plan["company_updates"]),
        "people_to_create": len(plan["person_creates"]),
        "people_to_update": len(plan["person_updates"]),
        "company_notes_to_create": len(plan["company_notes"]),
        "skipped_person_rows_without_name": len(plan["skipped_person_rows"]),
        "skipped_rows_with_invalid_company": len(plan["skipped_company_rows"]),
        "duplicate_source_emails": len(plan["source_duplicate_emails"]),
        "duplicate_source_email_examples": list(plan["source_duplicate_emails"].items())[:12],
        "company_create_examples": [action["data"].get("name") for action in plan["company_creates"][:10]],
        "company_update_examples": [action["id"] for action in plan["company_updates"][:10]],
    }


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Import Controlit CRM contacts from Excel files.")
    parser.add_argument("--dry-run", action="store_true", help="Plan only. This is the default unless --apply is passed.")
    parser.add_argument("--apply", action="store_true", help="Apply metadata/import changes.")
    parser.add_argument("--pilot", action="store_true", help="Only import pilot companies: YIT, Sensor Innovation, DayOne.")
    parser.add_argument("--limit-companies", type=int, default=None, help="Limit import to the first N source companies.")
    parser.add_argument("--setup-metadata-only", action="store_true", help="Only ensure Companies metadata fields/views.")
    parser.add_argument("--skip-backup", action="store_true", help="Skip DB backup for apply modes.")
    parser.add_argument("--somija-path", default=DEFAULT_SOMIJA_PATH)
    parser.add_argument("--crm-excel-path", default=DEFAULT_CRM_EXCEL_PATH)
    return parser.parse_args(argv)


def config_from_env() -> CrmConfig:
    email = os.environ.get("CRM_ADMIN_EMAIL")
    password = os.environ.get("CRM_ADMIN_PASSWORD")
    missing = [name for name, value in (("CRM_ADMIN_EMAIL", email), ("CRM_ADMIN_PASSWORD", password)) if not value]
    if missing:
        raise RuntimeError(f"Missing required env vars: {', '.join(missing)}")
    return CrmConfig(
        base_url=os.environ.get("CRM_BASE_URL", DEFAULT_BASE_URL).rstrip("/"),
        admin_email=email or "",
        admin_password=password or "",
        workspace_id=os.environ.get("CRM_WORKSPACE_ID", DEFAULT_WORKSPACE_ID),
        backup_dir=Path(os.environ.get("CRM_BACKUP_DIR", DEFAULT_BACKUP_DIR)).resolve(),
        ssh_host=os.environ.get("CRM_SSH_HOST", DEFAULT_SSH_HOST),
        remote_dir=os.environ.get("CRM_REMOTE_DIR", DEFAULT_REMOTE_DIR),
        db_service=os.environ.get("CRM_DB_SERVICE", DEFAULT_DB_SERVICE),
        db_user=os.environ.get("CRM_DB_USER", DEFAULT_DB_USER),
        db_name=os.environ.get("CRM_DB_NAME", DEFAULT_DB_NAME),
    )


def main(argv: list[str]) -> int:
    args = parse_args(argv)
    apply = args.apply
    config = config_from_env()
    client = CrmClient(config)
    client.authenticate()
    print(f"Workspace: {client.workspace['displayName']} ({client.workspace['id']})")

    if apply and not args.skip_backup:
        backup_path = create_database_backup(config, "controlit-crm-contact-import")
        print(f"Backup complete: {backup_path}")

    metadata_actions = ensure_company_metadata(client, apply=apply)
    if metadata_actions:
        print("Metadata changes:")
        for action in metadata_actions:
            print(f"- {action}")
    else:
        print("Metadata: no changes needed.")

    if args.setup_metadata_only:
        return 0

    source_records = parse_source_records(args.somija_path, args.crm_excel_path)
    filtered_records = filter_records(source_records, pilot=args.pilot, limit_companies=args.limit_companies)
    existing_state = fetch_existing_state(client)
    plan = build_import_plan(filtered_records, existing_state)
    summary = summarize_plan(plan)
    print(json.dumps(summary, ensure_ascii=False, indent=2))

    if not apply:
        print("Dry run complete. No CRM records were changed.")
        return 0

    applied = apply_import_plan(client, plan, existing_state)
    print(json.dumps({"applied": applied}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
